from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def iter_block_items(doc: Document):
    body = doc.element.body
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    for child in body.iterchildren():
        if child in paragraphs:
            yield "paragraph", paragraphs[child]
        elif child in tables:
            yield "table", tables[child]


def main() -> None:
    for raw_path in sys.argv[1:]:
        path = Path(raw_path)
        doc = Document(path)
        print(f"\n===== FILE: {path.name} =====")
        p_index = 0
        t_index = 0
        for kind, item in iter_block_items(doc):
            if kind == "paragraph":
                text = item.text.strip()
                if text:
                    p_index += 1
                    style = item.style.name if item.style else ""
                    print(f"[P{p_index}|{style}] {text}")
            else:
                t_index += 1
                print(f"[TABLE {t_index}]")
                for r_index, row in enumerate(item.rows, start=1):
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    print(f"[T{t_index}R{r_index}] " + " | ".join(cells))


if __name__ == "__main__":
    main()
